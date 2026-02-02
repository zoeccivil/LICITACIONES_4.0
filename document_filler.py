import os
import re
from docx import Document
from docx.oxml.ns import qn  # para usar QNames sin 'namespaces='

def fill_template(template_path, output_path, context, debug=False):
    """
    Reemplaza {{clave}} por su valor en TODO el .docx:
      - Cuerpo (párrafos y tablas, incl. anidadas)
      - Cuadros de texto (shapes, modernos y antiguos)
      - Encabezados y pies
    Soporta placeholders partidos entre 'runs' sin usar 'namespaces='.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"No se encontró la plantilla: {template_path}")

    doc = Document(template_path)

    # Diccionario de reemplazos: {{clave}} -> string(valor)
    repl = {f"{{{{{k}}}}}": ("" if v is None else str(v)) for k, v in (context or {}).items()}

    W_P = qn("w:p")
    W_T = qn("w:t")

    def _replace_in_paragraph_elm(p_elm):
        """
        p_elm es un elemento <w:p>. Unimos todos sus <w:t>, reemplazamos y
        escribimos el resultado en el primer <w:t>, vaciando los demás.
        """
        t_nodes = list(p_elm.iter(W_T))
        if not t_nodes:
            return 0
        original = "".join((t.text or "") for t in t_nodes)
        new_text = original
        for ph, val in repl.items():
            if ph in new_text:
                new_text = new_text.replace(ph, val)
        if new_text != original:
            t_nodes[0].text = new_text
            for t in t_nodes[1:]:
                t.text = ""
            return 1
        return 0

    def _process_root(root_elm):
        """Procesa TODOS los <w:p> bajo root_elm (incluye shapes y tablas)."""
        count = 0
        for p in root_elm.iter(W_P):
            count += _replace_in_paragraph_elm(p)
        return count

    total = 0

    # 1) Documento completo (cuerpo + tablas + shapes embebidos)
    total += _process_root(doc.element)

    # 2) Encabezados y pies (cada sección)
    for section in doc.sections:
        total += _process_root(section.header._element)
        total += _process_root(section.footer._element)

    if debug:
        print(f"[fill_template] Reemplazos realizados: {total}")
        # Reportar placeholders que queden
        pat = re.compile(r"\{\{[^}]+\}\}")
        restantes = set()
        def _collect(root_elm):
            for t in root_elm.iter(W_T):
                if t.text:
                    for m in pat.findall(t.text):
                        restantes.add(m)
        _collect(doc.element)
        for s in doc.sections:
            _collect(s.header._element)
            _collect(s.footer._element)
        if restantes:
            print("[fill_template] Placeholders NO reemplazados:", sorted(restantes))

    doc.save(output_path)
